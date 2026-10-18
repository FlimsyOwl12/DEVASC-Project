"""Build the editable course deck; optional tooling dependency: python-pptx."""

from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
BG = 'F5F2EA'
INK = '183447'
TEAL = '087F78'
WHITE = 'FFFFFF'
MUTED = '526574'

SLIDES = [
    ('cover', 'Routes you can read,\nchoose and save', [
        ('DEVASC Option 1', 'Python + GraphHopper REST API'),
        ('PalpakGPT', 'Activities 3 and 4 | October 18, 2026'),
    ], 'Introduce the project as an enhancement of the supplied lab. This is a team project; do not claim that the group wrote the routing engine. Jenkins success, SCM triggering, failure/recovery and final-main build #5 are verified. Instructor scope/polling acceptance and actual team assessment remain pending.'),
    ('cards', 'What problem are we solving?', [
        ('Choose the intended place', 'Several places may share the same name. The user needs enough detail to choose.'),
        ('Read a useful route', 'Show familiar units, clear totals and directions without asking users to read JSON.'),
        ('Keep the result', 'Save directions locally and reuse a route during the same session.'),
        ('Build on the lab', 'Retain the REST API workflow while improving input, output and error handling.'),
    ], 'Explain the customer problem before the implementation. The supplied lab uses GraphHopper; the instructor guide mentions MapQuest, so provider approval is still an explicit team action.'),
    ('cards', 'From terminal input to a reusable route', [
        ('1 / Choose options', 'Travel profile, metric/imperial units, summary/detailed output.'),
        ('2 / Geocode origin', 'Validate up to five candidates; show coordinates and accept the selected place.'),
        ('3 / Geocode destination', 'Choose the destination; stop safely if a location is not found.'),
        ('4 / Request routing', 'Send the selected coordinates and profile; validate the JSON path.'),
        ('5 / Display and remember', 'Print totals/directions; retain the last ten successful routes.'),
        ('6 / Export locally', 'TXT, JSON or CSV; check filenames and avoid silent overwrites.'),
    ], 'Trace one route through the actual functions. Geocoding and routing use HTTPS GET requests. History and exports reuse validated route data without making additional requests.'),
    ('cards', 'Core enhancements: F1 to F4', [
        ('F1 / Units', 'Metric or imperial values apply to totals and every direction step.'),
        ('F2 / Clear output', 'Aligned summary, HH:MM:SS duration and numbered plain-text directions.'),
        ('F3 / Output mode', 'Summary retains totals; detailed mode also shows directions.'),
        ('F4 / Response checks', 'Validate fields and numeric ranges; recover safely and redact the key.'),
    ], 'Show the command-line output in the demonstration. Point out that readable plain text works without color and that malformed replies do not need to become tracebacks or expose the key.'),
    ('code', 'F7 / Save a route. F8 / Choose a place.', [
        ('SYNTHETIC OFFLINE DEMO', '1. Central Park, Example City (lat: 10.0, lng: 20.0)\n2. Central Park, Example City (lat: 11.0, lng: 21.0)\nSelect a location: 2\n\nDistance    : 2.0 miles\nDuration    : 00:10:00\n1. Turn left onto Example Street (1.0 miles)\n2. Arrive at Demo Destination (1.0 miles)\n\nExport formats: TXT | JSON | CSV'),
    ], 'Run python scripts/demo_offline.py. The banner says SYNTHETIC: these are fixed mocked replies, not a new live GraphHopper route. The second coordinates must reach routing. Exports retain directions even when terminal summary mode is selected.'),
    ('cards', 'Menu and session history', [
        ('Find a Route', 'Run the profile, unit, output-mode and location-selection workflow.'),
        ('View Route History', 'Inspect any of the last ten successful routes without another HTTP request.'),
        ('Export Previous Route', 'Save an existing session record in another supported format.'),
        ('Session boundary', 'History is in memory. It disappears when the application exits.'),
    ], 'In the offline replay, show history viewing and TXT/CSV export after the initial JSON export. The script verifies that only the initial three mocked HTTP calls occurred. Explain that there is no persistent database.'),
    ('metrics', 'Verification we can reproduce', [
        ('69', 'offline unittest cases'),
        ('2', 'CI Python versions: 3.10 / 3.14'),
        ('0', 'live API calls in the offline suite'),
    ], 'Run python scripts/run_offline_tests.py. The runner clears the real key, disables dotenv loading and blocks unmocked HTTP/socket connections. Failure and empty discovery return nonzero status. Merged-main run: https://github.com/FlimsyOwl12/DEVASC-Project/actions/runs/39132132376 .'),
    ('cards', 'Live evidence is a separate record', [
        ('Car', '11.6 km | 1,079,659 ms | 10 direction steps'),
        ('Bike', '12.4 km | 2,532,665 ms | 39 direction steps'),
        ('Foot', '11.6 km | 8,585,308 ms | 72 direction steps'),
        ('Disconnected route', 'Manila to Honolulu: HTTP 400, connection between locations not found.'),
    ], 'These observations were recorded October 14 using Manila City Hall to Quezon Memorial Circle. They predate F7/F8, are historical account-specific results, and are not a claim that this presentation made live requests. Source: docs/live-verification.md. Never show the API key or full credential-bearing URL.'),
    ('cards', 'Collaboration and review evidence', [
        ('Scoped integration', 'F1/F2/F4/F3: PRs #21–#24. F9 pipeline code: PR #26.'),
        ('F7/F8 batch', 'PR #27 merged October 18; exports, candidate selection, menu and history.'),
        ('Review fixes', 'Show coordinates for identical labels; remove nine tracked cache files.'),
        ('Attribution', 'Confirm each member\'s real contribution. Upload identity is not sole authorship.'),
    ], 'The DRAFT records Justine Smith Rendaje as scrum leader, Wilhelm Ramesh Dimacisil as recorder, and John Jessienel Arquesola / Janwel Gilo as Prompt Engineers. Each member should describe work they actually did and can explain. Personal contribution/challenge/lesson statements are still awaiting confirmation. PR #27 has an LGTM comment recorded as COMMENTED, not a formal APPROVED review.'),
    ('cards', 'Activity 4: Jenkins is required', [
        ('Pipeline stages', 'Checkout → Install dependencies → Offline tests → Archive report.'),
        ('Automatic detection', 'SCM polling started original build #2 and Demo recovery #2. Webhook delivery is unverified.'),
        ('Failure and recovery', 'Demo #1: 70 tests, one intentional failure. Demo #2: 69 passing tests.'),
        ('Main verified', 'Build #5 at c6d34d5: 69 tests passed, report archived, SUCCESS.'),
    ], 'Jenkins 2.580.1 is running. Original #1 passed manually, and SCM polling started #2 automatically. On the separate demo job, #1 at bdd74ce failed exactly one intentional test; reverting only that test produced ed43bde and automatic Demo #2 passed 69 tests. After PRs #30/#28 merged, original build #5 checked out c6d34d5 and passed 69 tests in 0.486 seconds with its report archived. The original job targets main and the demo job is disabled with its history retained. Open the saved failure/recovery/main reports and screenshots. See docs/jenkins-evidence.md and handoff PR #32. Polling is verified; webhook delivery and instructor acceptance are not.'),
    ('cards', 'Limits and the final team handoff', [
        ('Provider limits', 'Coverage, quotas, profile access and availability depend on the API/account.'),
        ('Application limits', 'Terminal UI; no map/database; only ten in-memory session records.'),
        ('Team confirmation', 'Feature owners, actual contributions, scope approval and the presentation rehearsal.'),
        ('Assessment still open', 'Complete the official rubric and personal reflections from actual experience.'),
    ], 'Close by distinguishing implemented application features from unfinished course deliverables. The supplied instructor guide lists required rubrics but does not include the detailed student questions/scoring criteria. Obtain the actual rubric, rather than inventing grades or personal experience.'),
    ('sources', 'Evidence to open during questions', [
        ('Merged F7/F8 changes / PR #27', 'https://github.com/FlimsyOwl12/DEVASC-Project/pull/27'),
        ('Merged-main 69-test CI run', 'https://github.com/FlimsyOwl12/DEVASC-Project/actions/runs/39132132376'),
        ('Repository and task tracker', 'https://github.com/FlimsyOwl12/DEVASC-Project'),
        ('Original lab reference', 'https://github.com/Cocoasaur/devasc-graphhopper-api/tree/2289d4d'),
        ('Jenkins build record / local links inside', 'https://github.com/FlimsyOwl12/DEVASC-Project/blob/main/docs/jenkins-evidence.md'),
    ], 'Supporting local materials: docs/final-report.md and docs/test-results.md (final report PR); docs/demo-guide.md; docs/team-reflection.md; docs/jenkins-local-setup.md (Jenkins PR); docs/live-verification.md. Also use the supplied DEVASC Project.pdf and REST API lab. PRs #28/#30 are merged. Report PR #29, presentation PR #31 and shared handoff PR #32 still need review/merge. Build #5 verified main at c6d34d5; do not describe later documentation as already merged.'),
]


def color(value):
    return RGBColor.from_string(value)


def text(slide, value, x, y, width, height, size=20, fill=INK, bold=False, font='Arial', link=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color(fill)
        if link:
            run.hyperlink.address = link
    return box


def card(slide, x, y, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(WHITE)
    shape.line.fill.background()
    return shape


def build():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    deck.core_properties.title = 'DEVASC Option 1 — GraphHopper enhancements'
    deck.core_properties.author = 'PalpakGPT'
    deck.core_properties.last_modified_by = 'PalpakGPT'
    deck.core_properties.created = datetime(2026, 10, 18, tzinfo=timezone.utc)
    deck.core_properties.modified = datetime(2026, 10, 18, tzinfo=timezone.utc)
    notes = ['# Suggested speaker notes — Activity 3/4\n', 'Review contribution statements and show the recorded Jenkins failure, recovery and final-main reports/screenshots before presenting.\n']
    for number, (kind, title, entries, narration) in enumerate(SLIDES, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(BG)
        text(slide, 'DEVASC / OPTION 1', .65, .35, 10, .3, 12, TEAL, True)
        text(slide, f'{number:02d} / {len(SLIDES):02d}', 11.7, .35, 1, .3, 12, MUTED)
        text(slide, 'PalpakGPT  ·  Activities 3 + 4  ·  18 October 2026', .65, 7.05, 12, .25, 11, MUTED)
        slide.notes_slide.notes_text_frame.text = narration
        notes.append(f'## Slide {number} — {title.replace(chr(10), " ")}\n\n{narration}\n')
        if kind == 'cover':
            text(slide, title, .8, 1.6, 11.8, 2.1, 44, INK, True)
            text(slide, entries[0][1], .85, 4.1, 11, .6, 26, TEAL, True)
            text(slide, entries[1][0] + '\n' + entries[1][1], .85, 5.1, 11, .9, 20, MUTED)
            continue
        text(slide, title, .65, .9, 12, 1.05, 31, INK, True)
        if kind == 'cards':
            height = 1.35 if len(entries) > 4 else 2.05
            step = 1.5 if len(entries) > 4 else 2.3
            for index, (label, body) in enumerate(entries):
                x = .65 + (index % 2) * 6.15
                y = 2.05 + (index // 2) * step
                card(slide, x, y, 5.88, height)
                text(slide, label, x+.22, y+.16, 5.44, .42, 21, TEAL, True)
                text(slide, body, x+.22, y+.68, 5.4, height-.73, 17 if len(entries)>4 else 20)
        elif kind == 'code':
            shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.65), Inches(2.0), Inches(12), Inches(4.8))
            shape.fill.solid(); shape.fill.fore_color.rgb=color(INK); shape.line.fill.background()
            text(slide, entries[0][0], .95, 2.2, 11.4, .35, 13, '80D4C6', True)
            text(slide, entries[0][1], .95, 2.75, 11.4, 3.9, 17, WHITE, font='DejaVu Sans Mono')
        elif kind == 'metrics':
            for index, (value, label) in enumerate(entries):
                x=.65+index*4.08
                card(slide,x,2.3,3.82,3.2)
                text(slide,value,x+.25,2.7,3.3,1.2,64,TEAL,True)
                text(slide,label,x+.25,4.2,3.25,1.0,21)
            text(slide,'GitHub CI + Jenkins main build #5: all 69 tests passed.',.8,6.0,11.8,.55,22,MUTED)
        else:
            for index, (label, url) in enumerate(entries):
                y=2.1+index*1.0
                text(slide,label,.85,y,11.6,.5,24,TEAL,True,link=url)
                text(slide,url,.85,y+.52,11.6,.35,13,MUTED,link=url)
    deck.save(ROOT/'devasc-overview.pptx')
    (ROOT/'speaker-notes.md').write_text('\n'.join(notes))
    print(f'Built {len(deck.slides)} editable slides and speaker notes.')


if __name__ == '__main__':
    build()
